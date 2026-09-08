from pathlib import Path
from datetime import datetime
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from pptx import Presentation
from pptx.util import Inches, Pt
from db import get_bill, daily_sales

OUT=Path("generated")
OUT.mkdir(exist_ok=True)

def invoice_pdf(bill_id):
    data=get_bill(bill_id)
    if not data: raise ValueError("Bill not found")
    b=data["bill"]
    path=OUT/f"invoice_{bill_id}.pdf"
    c=canvas.Canvas(str(path),pagesize=A4)
    w,h=A4
    y=h-55
    c.setFont("Helvetica-Bold",18); c.drawString(45,y,"KiranaOps GST Invoice")
    y-=25; c.setFont("Helvetica",9)
    c.drawString(45,y,"Demo store • GSTIN 33ABCDE1234F1Z5")
    y-=25; c.drawString(45,y,f"Invoice #{bill_id}   {b['created_at']}")
    y-=25
    c.line(45,y,w-45,y); y-=20
    c.setFont("Helvetica-Bold",9)
    c.drawString(45,y,"Item"); c.drawString(260,y,"Qty"); c.drawString(320,y,"Rate"); c.drawString(390,y,"GST"); c.drawString(455,y,"Total")
    y-=15; c.setFont("Helvetica",9)
    for it in data["items"]:
        c.drawString(45,y,it["name"][:32])
        c.drawRightString(290,y,f"{it['qty']:g}")
        c.drawRightString(360,y,f"₹{it['unit_price']:.2f}")
        c.drawRightString(425,y,f"{it['line_gst']:.2f}")
        c.drawRightString(525,y,f"₹{it['line_subtotal']+it['line_gst']:.2f}")
        y-=17
        if y<70: c.showPage(); y=h-50
    y-=10; c.line(300,y,525,y); y-=18
    c.setFont("Helvetica-Bold",10)
    c.drawRightString(525,y,f"Taxable: ₹{b['total']:.2f}"); y-=17
    c.drawRightString(525,y,f"GST: ₹{b['gst_total']:.2f}"); y-=17
    c.drawRightString(525,y,f"Grand Total: ₹{b['total']+b['gst_total']:.2f}")
    y-=25; c.setFont("Helvetica",9)
    c.drawString(45,y,f"Payment: {b.get('payment_method') or 'Not specified'}")
    c.save()
    return str(path)

def analysis_pptx():
    s=daily_sales()
    path=OUT/f"sales_analysis_{datetime.now():%Y%m%d_%H%M%S}.pptx"
    prs=Presentation()
    slide=prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text="KiranaOps Sales Analysis"
    slide.placeholders[1].text=f"Generated {s['date']}"
    slide=prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text="Store snapshot"
    box=slide.shapes.add_textbox(Inches(1),Inches(1.5),Inches(8),Inches(3)).text_frame
    box.text=f"Sales: ₹{s['total']:.2f}\nGST collected: ₹{s['gst']:.2f}\nBills: {s['bills']}"
    p=box.add_paragraph(); p.text="Top items:"
    for x in s["top"]:
        p=box.add_paragraph(); p.text=f"{x['name']} — {x['qty']:g} units"
    prs.save(path)
    return str(path)
